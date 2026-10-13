from __future__ import annotations

import json
import os
import re
import shutil
import subprocess
import sys
import threading
import time
import uuid
from zipfile import BadZipFile, ZipFile
from datetime import datetime, timezone
from http import HTTPStatus
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.parse import urlparse


ROOT = Path(__file__).resolve().parents[2]
APP_DIR = ROOT / "app"
DATA_DIR = APP_DIR / "data"
JOBS_DIR = DATA_DIR / "jobs"
BACKUP_DIR = DATA_DIR / "backups"
YEAR_RUNS_DIR = ROOT / "outputs" / "year_runs"
COLLECTION_WORKBOOK = ROOT / "input" / "Diecast 2026.xlsx"
PYTHON = sys.executable
HOST = os.environ.get("F1_JOB_HOST", "127.0.0.1")
PORT = int(os.environ.get("F1_JOB_PORT", "8765"))
ADMIN_TOKEN = os.environ.get("F1_ADMIN_TOKEN", "")

JOBS_DIR.mkdir(parents=True, exist_ok=True)
BACKUP_DIR.mkdir(parents=True, exist_ok=True)

running_lock = threading.Lock()
running_jobs: dict[str, threading.Thread] = {}


def now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def clean_year(value: object) -> str:
    text = str(value or "").strip()
    if not re.fullmatch(r"(19[5-9]\d|20[0-3]\d)", text):
        raise ValueError("Rok musi byt v rozsahu 1950-2039.")
    return text


def job_paths(job_id: str) -> tuple[Path, Path]:
    return JOBS_DIR / f"{job_id}.json", JOBS_DIR / f"{job_id}.log"


def read_json(path: Path, fallback: object) -> object:
    if not path.exists():
        return fallback
    return json.loads(path.read_text(encoding="utf-8"))


def write_json(path: Path, payload: object) -> None:
    path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")


def append_log(job_id: str, message: str) -> None:
    _, log_path = job_paths(job_id)
    with log_path.open("a", encoding="utf-8", errors="replace") as handle:
        handle.write(message.rstrip() + "\n")


def load_job(job_id: str) -> dict:
    meta_path, _ = job_paths(job_id)
    payload = read_json(meta_path, {})
    if not isinstance(payload, dict):
        return {}
    return payload


def save_job(job: dict) -> None:
    meta_path, _ = job_paths(job["id"])
    write_json(meta_path, job)


def set_job_state(job: dict, status: str, step: str = "", error: str = "") -> None:
    job["status"] = status
    job["step"] = step
    job["error"] = error
    job["updatedAt"] = now_iso()
    if status in {"done", "failed"}:
        job["finishedAt"] = now_iso()
    save_job(job)


def run_command(job: dict, label: str, command: list[str]) -> None:
    append_log(job["id"], "")
    append_log(job["id"], f"== {label} ==")
    append_log(job["id"], " ".join(command))
    set_job_state(job, "running", label)
    env = os.environ.copy()
    env["PYTHONIOENCODING"] = "utf-8"
    process = subprocess.Popen(
        command,
        cwd=ROOT,
        env=env,
        stdout=subprocess.PIPE,
        stderr=subprocess.STDOUT,
        text=True,
        encoding="utf-8",
        errors="replace",
        bufsize=1,
    )
    assert process.stdout is not None
    for line in process.stdout:
        append_log(job["id"], line.rstrip())
    return_code = process.wait()
    if return_code:
        raise RuntimeError(f"Krok '{label}' skoncil chybou {return_code}.")


def backup_live_data(job: dict) -> None:
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    target = BACKUP_DIR / f"{job['season']}-{job['id']}-{stamp}"
    target.mkdir(parents=True, exist_ok=True)
    for name in ["app-data.json", "model_photo_overrides.json"]:
        source = DATA_DIR / name
        if source.exists():
            shutil.copy2(source, target / name)
    append_log(job["id"], f"Backup dat: {target}")


def summarize_year(season: str) -> dict:
    app_data = DATA_DIR / "app-data.json"
    if not app_data.exists():
        return {}
    data = json.loads(app_data.read_text(encoding="utf-8"))
    models = [row for row in data.get("models", []) if str(row.get("season")) == season]
    collection = [row for row in data.get("collection", []) if str(row.get("season")) == season]
    audit_path = YEAR_RUNS_DIR / str(season) / "photo_audit.json"
    problem_path = YEAR_RUNS_DIR / str(season) / "photo_audit_problem_rows.json"
    photo_status = {}
    photo_problems = []
    photo_by_manufacturer = {}
    if audit_path.exists():
        audit = json.loads(audit_path.read_text(encoding="utf-8"))
        if isinstance(audit, dict):
            photo_status = audit.get("byStatus", {})
            photo_by_manufacturer = audit.get("byManufacturer", {})
    if problem_path.exists():
        problem_rows = json.loads(problem_path.read_text(encoding="utf-8"))
        if isinstance(problem_rows, list):
            photo_problems = problem_rows[:80]
    with_photo = int(photo_status.get("verified", 0)) if photo_status else sum(
        1
        for row in models
        if row.get("mainPhoto") or any(str(url or "").strip() for url in row.get("photoUrls", []))
    )
    return {
        "models": len(models),
        "collectionRows": len(collection),
        "modelsWithPhoto": with_photo,
        "missingPhotos": max(len(models) - with_photo, 0),
        "photoStatus": photo_status,
        "photoByManufacturer": photo_by_manufacturer,
        "photoProblems": photo_problems,
    }


def run_year_job(job_id: str) -> None:
    job = load_job(job_id)
    try:
        season = job["season"]
        full_catalog = bool(job.get("fullCatalog", True))
        photo_limit = int(job.get("photoLimit", 250))
        append_log(job_id, f"Start rocniku {season}: {now_iso()}")
        backup_live_data(job)

        if full_catalog:
            run_command(job, "Kompletni sber katalogu ze zdroju", [PYTHON, "collect_model_catalog_expanded.py"])
        else:
            append_log(job_id, "Kompletni sber katalogu preskocen, pouzije se existujici master katalog.")

        run_command(job, "Import Minichamps z 143diecast cache", [PYTHON, "app/scripts/import_143diecast_minichamps.py"])
        run_command(job, "Import Spark official API", [PYTHON, "app/scripts/import_spark_official_f1.py", "--season", season])
        run_command(job, "Sestaveni pracovniho katalogu pro dohledani fotek", [PYTHON, "match_model_catalog.py"])
        run_command(job, "Sestaveni pracovni aplikacni kopie", [PYTHON, "app/scripts/prepare_app_data.py"])
        run_command(
            job,
            "Dohledani fotek pro rocnik",
            [PYTHON, "app/scripts/discover_model_photos.py", "--season", season, "--limit", str(photo_limit)],
        )
        run_command(job, "Finalni porovnani katalogu se sbirkou", [PYTHON, "match_model_catalog.py"])
        run_command(job, "Obnova dat aplikace po fotkach a porovnani", [PYTHON, "app/scripts/prepare_app_data.py"])
        run_command(job, "Audit realne dostupnosti fotek", [PYTHON, "app/scripts/audit_year_photos.py", "--season", season])

        job["summary"] = summarize_year(season)
        append_log(job_id, f"Hotovo: {json.dumps(job['summary'], ensure_ascii=False)}")
        set_job_state(job, "done", "Hotovo")
    except Exception as exc:
        append_log(job_id, f"CHYBA: {exc}")
        set_job_state(job, "failed", job.get("step", ""), str(exc))
    finally:
        with running_lock:
            running_jobs.pop(job_id, None)


def validate_collection_workbook(path: Path) -> int:
    try:
        with ZipFile(path) as archive:
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("Soubor není platný Excel XLSX.")
    except (BadZipFile, OSError) as exc:
        raise ValueError("Soubor není platný Excel XLSX.") from exc
    try:
        from openpyxl import load_workbook

        book = load_workbook(path, read_only=True, data_only=True)
        try:
            if "Overview" not in book.sheetnames:
                raise ValueError("V Excelu chybí list Overview.")
            sheet = book["Overview"]
            headers = [str(cell.value or "").strip() for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            required = {"Year", "Team", "Car", "Type", "Nr", "Driver", "Brand", "Pc", "Code"}
            missing = sorted(required - set(headers))
            if missing:
                raise ValueError(f"V listu Overview chybí sloupce: {', '.join(missing)}")
            return sum(1 for row in sheet.iter_rows(min_row=2, values_only=True) if any(value is not None and str(value).strip() for value in row))
        finally:
            book.close()
    except ImportError as exc:
        raise RuntimeError("Na serveru chybí knihovna openpyxl.") from exc


def run_collection_update(job_id: str, staged_path: Path, row_count: int) -> None:
    job = load_job(job_id)
    backup_dir = BACKUP_DIR / f"collection-{job_id}"
    targets = [
        COLLECTION_WORKBOOK,
        DATA_DIR / "app-data.json",
        ROOT / "outputs/model_catalog/all_f1_143_models_matched.json",
        ROOT / "outputs/model_catalog/all_f1_143_models_matched.csv",
        ROOT / "outputs/model_catalog/all_f1_143_models_summary.json",
    ]
    workbook_replaced = False
    try:
        backup_dir.mkdir(parents=True, exist_ok=True)
        job["step"] = "Záloha předchozí sbírky a dat"
        save_job(job)
        for target in targets:
            if target.exists():
                backup = backup_dir / target.name
                shutil.copy2(target, backup)
                (backup_dir / f"{target.name}.present").write_text("1", encoding="ascii")
        COLLECTION_WORKBOOK.parent.mkdir(parents=True, exist_ok=True)
        os.replace(staged_path, COLLECTION_WORKBOOK)
        workbook_replaced = True
        run_command(job, "Párování katalogu proti nové sbírce", [PYTHON, "match_model_catalog.py"])
        run_command(job, "Přepočet dat aplikace", [PYTHON, "app/scripts/prepare_app_data.py"])
        data = json.loads((DATA_DIR / "app-data.json").read_text(encoding="utf-8"))
        job["summary"] = {
            "sourceRows": row_count,
            "collectionRows": len(data.get("collectionItems", [])),
            "masterModels": len(data.get("models", [])),
            "updatedAt": data.get("generatedAt", ""),
        }
        append_log(job_id, f"Sbírka načtena: {json.dumps(job['summary'], ensure_ascii=False)}")
        set_job_state(job, "done", "Sbírka načtena")
    except Exception as exc:
        append_log(job_id, f"CHYBA: {exc}; obnovuji předchozí data.")
        if workbook_replaced:
            for target in targets:
                backup = backup_dir / target.name
                if (backup_dir / f"{target.name}.present").exists():
                    if backup.exists():
                        shutil.copy2(backup, target)
                elif target.exists():
                    target.unlink()
        set_job_state(job, "failed", job.get("step", ""), str(exc))
    finally:
        staged_path.unlink(missing_ok=True)
        with running_lock:
            running_jobs.pop(job_id, None)


def list_jobs() -> list[dict]:
    jobs = []
    for path in sorted(JOBS_DIR.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
        payload = read_json(path, {})
        if isinstance(payload, dict):
            jobs.append(payload)
    return jobs[:30]


class Handler(BaseHTTPRequestHandler):
    server_version = "F1JobServer/1.0"

    def log_message(self, format: str, *args: object) -> None:
        return

    def send_payload(self, status: int, payload: object, content_type: str = "application/json") -> None:
        if content_type == "application/json":
            body = json.dumps(payload, ensure_ascii=False).encode("utf-8")
        else:
            body = str(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", f"{content_type}; charset=utf-8")
        self.send_header("Cache-Control", "no-store")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_body(self) -> dict:
        length = int(self.headers.get("Content-Length", "0") or "0")
        if not length:
            return {}
        raw = self.rfile.read(length).decode("utf-8")
        payload = json.loads(raw)
        if not isinstance(payload, dict):
            raise ValueError("Body musi byt JSON objekt.")
        return payload

    def require_admin(self) -> bool:
        if not ADMIN_TOKEN:
            self.send_payload(HTTPStatus.FORBIDDEN, {"error": "Na serveru neni nastaveny F1_ADMIN_TOKEN."})
            return False
        supplied = self.headers.get("X-Admin-Token", "")
        if supplied != ADMIN_TOKEN:
            self.send_payload(HTTPStatus.UNAUTHORIZED, {"error": "Neplatny admin token."})
            return False
        return True

    def do_GET(self) -> None:
        path = urlparse(self.path).path
        if path == "/api/health":
            self.send_payload(HTTPStatus.OK, {"ok": True, "time": now_iso()})
            return
        if path == "/api/jobs":
            self.send_payload(HTTPStatus.OK, {"jobs": list_jobs()})
            return
        match = re.fullmatch(r"/api/jobs/([a-zA-Z0-9-]+)", path)
        if match:
            job = load_job(match.group(1))
            if not job:
                self.send_payload(HTTPStatus.NOT_FOUND, {"error": "Uloha nenalezena."})
                return
            self.send_payload(HTTPStatus.OK, {"job": job})
            return
        match = re.fullmatch(r"/api/jobs/([a-zA-Z0-9-]+)/log", path)
        if match:
            _, log_path = job_paths(match.group(1))
            if not log_path.exists():
                self.send_payload(HTTPStatus.NOT_FOUND, "Log nenalezen.", "text/plain")
                return
            self.send_payload(HTTPStatus.OK, log_path.read_text(encoding="utf-8", errors="replace"), "text/plain")
            return
        self.send_payload(HTTPStatus.NOT_FOUND, {"error": "Nenalezeno."})

    def do_POST(self) -> None:
        path = urlparse(self.path).path
        if path not in {"/api/jobs/start", "/api/collection/update"}:
            self.send_payload(HTTPStatus.NOT_FOUND, {"error": "Nenalezeno."})
            return
        if not self.require_admin():
            return
        try:
            with running_lock:
                if any(thread.is_alive() for thread in running_jobs.values()):
                    self.send_payload(HTTPStatus.CONFLICT, {"error": "Jiná úloha právě běží. Počkejte na její dokončení."})
                    return
            if path == "/api/collection/update":
                length = int(self.headers.get("Content-Length", "0") or "0")
                if length <= 0 or length > 15 * 1024 * 1024:
                    raise ValueError("Soubor musí mít velikost do 15 MB.")
                raw = self.rfile.read(length)
                staged_path = DATA_DIR / f"collection-upload-{uuid.uuid4().hex}.xlsx"
                staged_path.write_bytes(raw)
                try:
                    row_count = validate_collection_workbook(staged_path)
                except Exception:
                    staged_path.unlink(missing_ok=True)
                    raise
                job_id = uuid.uuid4().hex[:12]
                job = {
                    "id": job_id,
                    "kind": "collection_update",
                    "season": "all",
                    "status": "queued",
                    "step": "Čeká na spuštění",
                    "error": "",
                    "createdAt": now_iso(),
                    "updatedAt": now_iso(),
                    "finishedAt": "",
                    "summary": {"sourceRows": row_count},
                }
                save_job(job)
                append_log(job_id, f"Aktualizace zdrojové sbírky ({row_count} datových řádků).")
                thread = threading.Thread(target=run_collection_update, args=(job_id, staged_path, row_count), daemon=True)
                with running_lock:
                    running_jobs[job_id] = thread
                thread.start()
                self.send_payload(HTTPStatus.ACCEPTED, {"job": job})
                return
            payload = self.read_body()
            season = clean_year(payload.get("season"))
            job_id = uuid.uuid4().hex[:12]
            job = {
                "id": job_id,
                "season": season,
                "status": "queued",
                "step": "Ceka na spusteni",
                "error": "",
                "createdAt": now_iso(),
                "updatedAt": now_iso(),
                "finishedAt": "",
                "fullCatalog": bool(payload.get("fullCatalog", True)),
                "photoLimit": int(payload.get("photoLimit", 250)),
                "summary": {},
            }
            save_job(job)
            append_log(job_id, f"Uloha vytvorena pro rok {season}.")
            thread = threading.Thread(target=run_year_job, args=(job_id,), daemon=True)
            with running_lock:
                running_jobs[job_id] = thread
            thread.start()
            self.send_payload(HTTPStatus.ACCEPTED, {"job": job})
        except Exception as exc:
            self.send_payload(HTTPStatus.BAD_REQUEST, {"error": str(exc)})


def main() -> None:
    print(f"F1 job server listening on {HOST}:{PORT}", flush=True)
    ThreadingHTTPServer((HOST, PORT), Handler).serve_forever()


if __name__ == "__main__":
    main()
